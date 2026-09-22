import openpyxl
import os
import json
import re

def parse_option_field(field_str, idx):
    field_str = field_str.strip()
    if not field_str:
        return None
        
    m = re.match(r"^([^\[]+)\[(.*)\]$", field_str)
    if m:
        label = m.group(1).strip()
        opts_raw = m.group(2).strip()
    else:
        label = field_str
        opts_raw = ""
        
    field_id = f"f_{idx}"
    is_support = "지주" in label
    is_photo = "사진" in label
    
    if opts_raw:
        if opts_raw.startswith("숫자:"):
            def_val = opts_raw.split(":", 1)[1].strip()
            return {
                "id": field_id,
                "label": label,
                "type": "number",
                "default": def_val,
                "isSupport": is_support,
                "isPhoto": is_photo
            }
        elif opts_raw == "직접입력" or opts_raw == "텍스트":
            return {
                "id": field_id,
                "label": label,
                "type": "text",
                "default": "",
                "isSupport": is_support,
                "isPhoto": is_photo
            }
        else:
            opts = [o.strip() for o in opts_raw.split(",") if o.strip()]
            def_val = opts[0] if opts else ""
            return {
                "id": field_id,
                "label": label,
                "type": "select",
                "options": opts,
                "default": def_val,
                "isSupport": is_support,
                "isPhoto": is_photo
            }
    else:
        return {
            "id": field_id,
            "label": label,
            "type": "text",
            "default": "",
            "isSupport": is_support,
            "isPhoto": is_photo
        }

def convert_excel_to_config(excel_path, output_js_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["시설물_조사옵션"]
    
    facility_config = {}
    
    # 5행부터 데이터 행
    for r in range(5, ws.max_row + 1):
        cell_a = ws.cell(row=r, column=1) # 시설물명
        cell_b = ws.cell(row=r, column=2) # 감지레이어
        cell_c = ws.cell(row=r, column=3) # 캐드레이어
        cell_d = ws.cell(row=r, column=4) # 색상
        cell_e = ws.cell(row=r, column=5) # 접두어
        
        name = cell_a.value
        if not name:
            continue
        name = str(name).strip()
        
        det_layer_str = str(cell_b.value or "").strip()
        det_layers = [l.strip() for l in det_layer_str.split(",") if l.strip()] if det_layer_str else []
        
        # B열 글자색이 파란색 계열인지 확인
        font_b = cell_b.font
        is_sub_attachable = False
        if font_b and font_b.color:
            rgb = getattr(font_b.color, 'rgb', '')
            if rgb and any(b_code in str(rgb).upper() for b_code in ["0070C0", "0000FF", "2563EB", "1E40AF", "3B82F6"]):
                is_sub_attachable = True
                
        cad_layer = str(cell_c.value or "").strip()
        if not cad_layer or cad_layer == "None":
            cad_layer = name + "_T"
            
        color_val = str(cell_d.value or "").strip()
        color_match = re.search(r"\d+", color_val)
        color_num = int(color_match.group()) if color_match else 7
        
        prefix = str(cell_e.value or "").strip() if cell_e.value else name
        
        fields = []
        field_idx = 1
        for col in range(6, ws.max_column + 1):
            cell_val = ws.cell(row=r, column=col).value
            if cell_val is not None:
                val_str = str(cell_val).strip()
                if val_str:
                    parsed_field = parse_option_field(val_str, field_idx)
                    if parsed_field:
                        fields.append(parsed_field)
                        field_idx += 1
                        
        facility_config[name] = {
            "title": name,
            "layer": cad_layer,
            "color": color_num,
            "prefix": prefix,
            "detectionLayers": det_layers,
            "isSubAttachable": is_sub_attachable,
            "fields": fields
        }
        
    js_content = "/**\n * 도로대장 시설물 조사 옵션 설정 (엑셀 기반 자동 생성)\n */\n"
    js_content += "(function (global) {\n  'use strict';\n\n"
    js_content += "  global.FACILITY_CONFIG = " + json.dumps(facility_config, ensure_ascii=False, indent=2) + ";\n"
    js_content += "})(typeof window !== 'undefined' ? window : this);\n"
    
    with open(output_js_path, "w", encoding="utf-8") as f:
        f.write(js_content)
        
    print(f"SUCCESS: Converted {len(facility_config)} facilities to {output_js_path}")

if __name__ == "__main__":
    excel_path = r"d:\!!프로그램\CODING\NDMAP_edit\PDMAP_평택용도로대장\20260918\도로대장_시설물_조사옵션_템플릿.xlsx"
    out_js = r"d:\!!프로그램\CODING\NDMAP_edit\PDMAP_평택용도로대장\20260918\facility-config.js"
    convert_excel_to_config(excel_path, out_js)
